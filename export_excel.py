"""Reference-format Excel export for Producer Calendar local app.

The exported workbook uses the SPE block-calendar workbook the user provided as
its visual template. Visible output tabs are one tab per schedule year, while an
Inputs tab lets the user continue changing dates/durations after export. The
year tabs remain formula-driven: month grids, top project lines, lower Dates box,
week counters, and color highlighting update from workbook formulas and
conditional formatting.
"""
from __future__ import annotations

from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import MergedCell

from calendar_engine import (
    calculate_schedule,
    parse_date,
    fmt,
    holidays_for_year,
    PERIOD_LABELS,
    PERIOD_COLORS,
    LOCATION_LABELS,
)

APP_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = APP_DIR / "templates" / "SPEBlockCalendar_Template.xlsx"

# The user asked for the reference output with no Vertex42 or Cleve/Movie Tools branding.
BANNED_PHRASES = ["Vertex42", "Yearly Calendar Template", "Movie Tools", "zcleve.com", "Cleve"]

INPUT_ROWS = {
    "rd": 9,
    "pre": 10,
    "travel": 11,
    "production": 12,
    "hiatus": 13,
    "post": 14,
    "print_ship": 15,
}

SCHEDULE_ROWS = {
    "rd": 2,
    "pre": 3,
    "travel": 4,
    "production": 5,
    "hiatus": 6,
    "post": 7,
    "print_ship": 8,
    "ready": 9,
}

# Helper rows written into hidden columns on each year sheet. These local cells
# make conditional formatting compatible with Excel, because CF formulas do not
# need to reference another worksheet directly.
HELPER_START_ROW = 2
HELPER_COLS = {
    "key": "AJ",
    "label": "AK",
    "start": "AL",
    "end": "AM",
    "metric": "AN",
    "color": "AO",
    "used": "AP",
    "holiday_date": "AQ",
    "holiday_name": "AR",
}

DATE_BLOCKS: List[Tuple[str, str, str]] = [
    ("A", "B", "H"), ("I", "J", "P"), ("Q", "R", "X"), ("Y", "Z", "AF"),
]
DATE_ROW_STARTS = [13, 22, 31]


def _set_reference_column_widths(ws: Worksheet) -> None:
    """Keep the reference calendar columns stable across exports.

    Excel column widths are global for the sheet. Earlier versions adjusted
    columns inside the lower Dates/Legend block, which made the calendar
    drift wider toward the right side of the sheet. This pins every calendar
    date column to the same width, using the B:H block as the visual baseline,
    and leaves the lower legend to work within those same columns.
    """
    date_width = 5.1
    gutter_width = 3.1
    # A, I, Q, and Y are the narrow week-number/gutter columns before each
    # four-month block. All seven day columns in every month block are equal.
    gutter_cols = {1, 9, 17, 25}
    for col_idx in range(1, 34):  # A:AG visible output area
        letter = get_column_letter(col_idx)
        if col_idx in gutter_cols:
            ws.column_dimensions[letter].width = gutter_width
        else:
            ws.column_dimensions[letter].width = date_width


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.replace("#", ""))


def _font(color: str = "000000", bold: bool = False, size: int = 10, name: str = "Arial") -> Font:
    return Font(name=name, size=size, bold=bold, color=color.replace("#", ""))


def _border(color: str = "B7B7B7") -> Border:
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def _date_value(value: Any) -> Optional[date]:
    return parse_date(value)


def _duration_from_payload(payload: Dict[str, Any], key: str, field: str, default: int = 0) -> int:
    try:
        return int(((payload.get("periods") or {}).get(key) or {}).get(field) or default)
    except Exception:
        return default


def _record_by_key(schedule: Dict[str, Any]) -> Dict[str, Dict[str, Any]]:
    return {p.get("key"): p for p in schedule.get("periods", [])}


def _set_date(cell, value: Optional[date]) -> None:
    cell.value = value
    cell.number_format = "m/d/yy"


def _clean_branding(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(phrase.lower() in cell.value.lower() for phrase in BANNED_PHRASES):
                    cell.value = None


def _remove_unwanted_sheets(wb, keep_years: List[int]) -> None:
    keep_names = {str(y) for y in keep_years}
    keep_names.update({"Inputs", "Schedule Data", "Holidays"})
    for ws in list(wb.worksheets):
        if ws.title not in keep_names:
            wb.remove(ws)


def _month_title_formula(year: int) -> str:
    return f'=TEXT(DATE({year},1,1),"mmmm")'


def _make_inputs_sheet(wb, schedule: Dict[str, Any], payload: Dict[str, Any]) -> Worksheet:
    if "Inputs" in wb.sheetnames:
        del wb["Inputs"]
    ws = wb.create_sheet("Inputs", 0)
    ws.sheet_view.showGridLines = False

    title = schedule.get("projectTitle") or "Feature Film"
    location_code = schedule.get("productionLocation") or "US"
    location_label = schedule.get("productionLocationLabel") or LOCATION_LABELS.get(location_code, "United States")
    recs = _record_by_key(schedule)
    ready = schedule.get("ready") or {}
    ready_date = _date_value(ready.get("date"))

    ws["A1"] = "Producer Calendar Inputs"
    ws["A1"].font = _font("FFFFFF", True, 16)
    ws["A1"].fill = _fill("1F4E78")
    ws.merge_cells("A1:H1")

    as_of_date = _date_value(payload.get("asOfDate") or payload.get("as_of") or payload.get("asOf")) or date.today()
    meta = [
        ("Project Title", title),
        ("Production Location", location_label),
        ("As Of Date", as_of_date),
        ("Scheduling Anchor Used By App", schedule.get("anchorLabel") or ""),
        ("Ready for Release Override", ready_date),
        ("Instructions", "Edit the blue input cells below. Week-based phases count Monday-Friday workweeks. Year tabs update by formulas/conditional formatting."),
    ]
    for r, (label, value) in enumerate(meta, 2):
        ws.cell(r, 1, label)
        ws.cell(r, 1).font = _font("000000", True, 10)
        ws.cell(r, 1).fill = _fill("D9EAF7")
        ws.cell(r, 2, value)
        if isinstance(value, date):
            ws.cell(r, 2).number_format = "m/d/yy"
        ws.cell(r, 1).border = _border()
        ws.cell(r, 2).border = _border()
        ws.cell(r, 2).fill = _fill("EAF3FF")
        ws.cell(r, 2).alignment = Alignment(wrap_text=True)

    headers = ["Period", "Start Date", "Weeks", "Production Days", "End Date", "Color", "Notes"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(8, c, h)
        cell.fill = _fill("1F4E78")
        cell.font = _font("FFFFFF", True, 10)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _border("FFFFFF")

    def p_start(key: str) -> Optional[date]:
        return _date_value((recs.get(key) or {}).get("start"))

    def p_end(key: str) -> Optional[date]:
        return _date_value((recs.get(key) or {}).get("end"))

    period_rows = [
        ("rd", "R&D", _duration_from_payload(payload, "rd", "weeks", 0), None),
        ("pre", "Pre-Production", _duration_from_payload(payload, "pre", "weeks", 0), None),
        ("travel", "Travel/Prep", _duration_from_payload(payload, "travel", "weeks", 0), None),
        ("production", "Production", None, _duration_from_payload(payload, "production", "days", 0)),
        ("hiatus", "Hiatus", None, None),
        ("post", "Post Production", _duration_from_payload(payload, "post", "weeks", 0), None),
        ("print_ship", "Print & Ship", _duration_from_payload(payload, "print_ship", "weeks", 0), None),
    ]
    for key, label, weeks, prod_days in period_rows:
        r = INPUT_ROWS[key]
        ws.cell(r, 1, label)
        _set_date(ws.cell(r, 2), p_start(key))
        if weeks is not None:
            ws.cell(r, 3, weeks)
        if prod_days is not None:
            ws.cell(r, 4, prod_days)
        if key == "hiatus":
            _set_date(ws.cell(r, 5), p_end(key))
        elif key == "production":
            # Production is weekdays only; selected-location holidays on the Holidays sheet are excluded.
            ws.cell(r, 5, f'=IF(AND(B{r}<>"",D{r}>0),WORKDAY(B{r}-1,D{r},Holidays!$A$4:$A$200),"")')
            ws.cell(r, 5).number_format = "m/d/yy"
        else:
            # All week-based phases count Monday-Friday workweeks. Holidays only extend Production.
            ws.cell(r, 5, f'=IF(AND(B{r}<>"",C{r}>0),WORKDAY(B{r}-1,C{r}*5),"")')
            ws.cell(r, 5).number_format = "m/d/yy"
        ws.cell(r, 6, PERIOD_COLORS[key])
        ws.cell(r, 7, "Editable start/duration. Week-based ends use Monday-Friday workweeks." if key != "hiatus" else "Single editable hiatus date range.")
        fill = _fill(PERIOD_COLORS[key])
        for c in range(1, 8):
            ws.cell(r, c).border = _border()
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="center")
            if c == 1:
                ws.cell(r, c).fill = fill
                if key == "ready":
                    ws.cell(r, c).font = _font("FFFFFF", True)
        # Blue input cells.
        for c in [2, 3, 4, 5] if key == "hiatus" else [2, 3, 4]:
            ws.cell(r, c).fill = _fill("EAF3FF")

    widths = {"A": 24, "B": 14, "C": 10, "D": 16, "E": 14, "F": 12, "G": 44, "H": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A8"
    return ws


def _make_holidays_sheet(wb, schedule: Dict[str, Any], years: List[int]) -> Worksheet:
    if "Holidays" in wb.sheetnames:
        del wb["Holidays"]
    ws = wb.create_sheet("Holidays")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Selected Production Location Holidays"
    ws["A1"].font = _font("FFFFFF", True, 14)
    ws["A1"].fill = _fill("1F4E78")
    ws.merge_cells("A1:D1")
    headers = ["Date", "Holiday", "Region", "Observed"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, h).fill = _fill("D9EAF7")
        ws.cell(3, c).font = _font("000000", True)
        ws.cell(3, c).border = _border()
    region = schedule.get("productionLocation") or "US"
    row = 4
    # Include a one-year buffer so production formulas still work if a user edits
    # a start/end date near the edge of the schedule years.
    for year in range(min(years) - 1, max(years) + 2):
        for h in holidays_for_year(year, region):
            ws.cell(row, 1, h.date)
            ws.cell(row, 1).number_format = "m/d/yy"
            ws.cell(row, 2, h.name)
            ws.cell(row, 3, h.region)
            ws.cell(row, 4, "Yes" if h.observed else "")
            for c in range(1, 5):
                ws.cell(row, c).border = _border()
            row += 1
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 12
    return ws


def _make_schedule_data_sheet(wb, schedule: Dict[str, Any], payload: Dict[str, Any]) -> Worksheet:
    if "Schedule Data" in wb.sheetnames:
        del wb["Schedule Data"]
    ws = wb.create_sheet("Schedule Data")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Schedule Data - formula source for year tabs"
    ws["A1"].font = _font("FFFFFF", True, 12)
    ws["A1"].fill = _fill("1F4E78")
    ws.merge_cells("A1:L1")
    headers = ["Key", "Period", "Start", "End", "Metric", "Metric Type", "Color", "Used", "Start Text", "Date Box Line", "Metric Text", "Notes"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(2, c, h)
        cell.fill = _fill("D9EAF7")
        cell.font = _font("000000", True)
        cell.border = _border()
    # Shift period data to rows 3:10 so row 2 can be headers.
    data_rows = {
        "rd": 3, "pre": 4, "travel": 5, "production": 6, "hiatus": 7, "post": 8, "print_ship": 9, "ready": 10
    }
    input_for_key = INPUT_ROWS
    key_order = ["rd", "pre", "travel", "production", "hiatus", "post", "print_ship", "ready"]
    for key in key_order:
        r = data_rows[key]
        ws.cell(r, 1, key)
        ws.cell(r, 2, PERIOD_LABELS[key])
        if key == "ready":
            # If an override is entered, use it; otherwise calculate release as the day after Print & Ship.
            ws.cell(r, 3, '=IF(Inputs!$B$6<>"",Inputs!$B$6,IF(Inputs!$E$15<>"",Inputs!$E$15+1,""))')
            ws.cell(r, 4, f"=C{r}")
            ws.cell(r, 5, "")
            ws.cell(r, 6, "single date")
        else:
            ir = input_for_key[key]
            ws.cell(r, 3, f"=Inputs!$B${ir}")
            ws.cell(r, 4, f"=Inputs!$E${ir}")
            if key == "production":
                ws.cell(r, 5, f"=Inputs!$D${ir}")
                ws.cell(r, 6, "workdays")
            elif key == "hiatus":
                ws.cell(r, 5, f'=IF(AND(C{r}<>"",D{r}<>""),D{r}-C{r}+1,"")')
                ws.cell(r, 6, "date range")
            else:
                ws.cell(r, 5, f"=Inputs!$C${ir}")
                ws.cell(r, 6, "weeks")
        ws.cell(r, 7, PERIOD_COLORS[key])
        if key == "ready":
            ws.cell(r, 8, f'=C{r}<>""')
        else:
            ws.cell(r, 8, f'=AND(C{r}<>"",D{r}<>"",D{r}>=C{r})')
        # Ordinal date text formula.
        suffix_c = f'IF(AND(MOD(DAY(C{r}),100)>=11,MOD(DAY(C{r}),100)<=13),"th",IF(MOD(DAY(C{r}),10)=1,"st",IF(MOD(DAY(C{r}),10)=2,"nd",IF(MOD(DAY(C{r}),10)=3,"rd","th"))))'
        ws.cell(r, 9, f'=IF(C{r}="","",TEXT(C{r},"mmmm d")&{suffix_c}&", "&TEXT(C{r},"yyyy"))')
        if key == "print_ship":
            ws.cell(r, 10, f'=IF(H{r},"Plus "&E{r}&" weeks for Print & Ship","")')
        elif key == "ready":
            ws.cell(r, 10, f'=IF(H{r},"Ready for Release - "&I{r},"")')
        elif key == "hiatus":
            suffix_d = f'IF(AND(MOD(DAY(D{r}),100)>=11,MOD(DAY(D{r}),100)<=13),"th",IF(MOD(DAY(D{r}),10)=1,"st",IF(MOD(DAY(D{r}),10)=2,"nd",IF(MOD(DAY(D{r}),10)=3,"rd","th"))))'
            ws.cell(r, 10, f'=IF(H{r},"Hiatus - "&I{r}&" to "&TEXT(D{r},"mmmm d")&{suffix_d}&", "&TEXT(D{r},"yyyy"),"")')
        else:
            ws.cell(r, 10, f'=IF(H{r},"Begin "&B{r}&" - "&I{r},"")')
        if key == "production":
            ws.cell(r, 11, f'=IF(H{r},ROUNDUP((D{r}-C{r}+1)/7,0)&" weeks /"&CHAR(10)&E{r}&" days","")')
        elif key == "hiatus" or key == "ready":
            ws.cell(r, 11, "")
        else:
            ws.cell(r, 11, f'=IF(H{r},E{r}&" weeks","")')
        for c in [3, 4]:
            ws.cell(r, c).number_format = "m/d/yy"
        for c in range(1, 12):
            ws.cell(r, c).border = _border()
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in {"A": 12, "B": 20, "C": 13, "D": 13, "E": 10, "F": 12, "G": 12, "H": 9, "I": 26, "J": 58, "K": 18}.items():
        ws.column_dimensions[col].width = width
    return ws


def _prepare_year_sheet(ws: Worksheet, year: int) -> None:
    ws.title = str(year)
    ws.sheet_view.showGridLines = False
    _set_reference_column_widths(ws)
    # Remove any leftover template logos/images from the reference workbook.
    try:
        ws._images = []
    except Exception:
        pass
    ws["D3"] = year
    ws["J3"] = 1
    ws["R3"] = 1
    # Remove external template/vendor content while preserving formula controls.
    for cell in ["A1", "R1", "AF3", "AG3", "AH3", "AI3", "AI12"]:
        ws[cell] = None
    for row in range(1, 7):
        ws.row_dimensions[row].hidden = True
    for col in ["AH", "AI"]:
        ws.column_dimensions[col].hidden = True
    # Hidden helper columns used by formulas/conditional formatting.
    for col_idx in range(36, 45):  # AJ:AR
        ws.column_dimensions[get_column_letter(col_idx)].hidden = True
    ws.print_area = "A7:AG48"
    ws.page_setup.orientation = "landscape"
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.35
    try:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
    except Exception:
        pass


def _populate_year_headers(ws: Worksheet, year: int) -> None:
    ws["B7"] = '=Inputs!$B$2&", Ready for Release "&TEXT(\'Schedule Data\'!$C$10,"mmmm d, yyyy")&" as of "&TEXT(Inputs!$B$4,"m/d/yy")'
    ws["B8"] = '="Assumes "&Inputs!$D$12&" Shoot Days; "&Inputs!$C$14&" Workweek Post"'
    ws["B9"] = "JANUARY - DECEMBER"
    ws["G9"] = f"={year}"
    ws["B7"].font = _font("000000", True, 13)
    ws["B8"].font = _font("000000", True, 10)


def _populate_local_helpers(ws: Worksheet, year: int, schedule: Dict[str, Any]) -> None:
    schedule_row_map = {"rd": 3, "pre": 4, "travel": 5, "production": 6, "hiatus": 7, "post": 8, "print_ship": 9, "ready": 10}
    for idx, key in enumerate(["rd", "pre", "travel", "production", "hiatus", "post", "print_ship", "ready"], HELPER_START_ROW):
        ws[f"{HELPER_COLS['key']}{idx}"] = key
        ws[f"{HELPER_COLS['label']}{idx}"] = PERIOD_LABELS[key]
        sd = schedule_row_map[key]
        ws[f"{HELPER_COLS['start']}{idx}"] = f"='Schedule Data'!$C${sd}"
        ws[f"{HELPER_COLS['end']}{idx}"] = f"='Schedule Data'!$D${sd}"
        ws[f"{HELPER_COLS['metric']}{idx}"] = f"='Schedule Data'!$E${sd}"
        ws[f"{HELPER_COLS['color']}{idx}"] = PERIOD_COLORS[key]
        ws[f"{HELPER_COLS['used']}{idx}"] = f"='Schedule Data'!$H${sd}"
        ws[f"{HELPER_COLS['start']}{idx}"].number_format = "m/d/yy"
        ws[f"{HELPER_COLS['end']}{idx}"].number_format = "m/d/yy"

    # Holiday helper list for this year. Conditional formatting highlights these
    # only when they overlap the Production date range.
    region = schedule.get("productionLocation") or "US"
    hrow = 2
    ws[f"{HELPER_COLS['holiday_date']}{hrow}"] = "Holiday Date"
    ws[f"{HELPER_COLS['holiday_name']}{hrow}"] = "Holiday"
    hrow += 1
    for h in holidays_for_year(year, region):
        ws[f"{HELPER_COLS['holiday_date']}{hrow}"] = h.date
        ws[f"{HELPER_COLS['holiday_date']}{hrow}"].number_format = "m/d/yy"
        ws[f"{HELPER_COLS['holiday_name']}{hrow}"] = h.name
        hrow += 1


def _all_date_ranges() -> List[str]:
    ranges = []
    for row_start in DATE_ROW_STARTS:
        for _, start_col, end_col in DATE_BLOCKS:
            ranges.append(f"{start_col}{row_start}:{end_col}{row_start+5}")
    return ranges


def _clear_conditional_formatting(ws: Worksheet) -> None:
    try:
        ws.conditional_formatting._cf_rules.clear()
    except Exception:
        pass


def _add_cf(ws: Worksheet, ranges: List[str], formula: str, fill: Optional[str] = None, font_color: Optional[str] = None, bold: bool = False, stop: bool = False) -> None:
    kwargs: Dict[str, Any] = {"formula": [formula], "stopIfTrue": stop}
    if fill:
        kwargs["fill"] = _fill(fill)
    if font_color or bold:
        kwargs["font"] = _font(font_color or "000000", bold, 10)
    rule = FormulaRule(**kwargs)
    for rng in ranges:
        ws.conditional_formatting.add(rng, rule)


def _apply_calendar_conditional_formatting(ws: Worksheet) -> None:
    _clear_conditional_formatting(ws)
    all_ranges = _all_date_ranges()
    combined = " ".join(all_ranges)
    # Ready for Release is the highest-priority milestone.
    _add_cf(
        ws,
        [combined],
        '=AND(ISNUMBER(B13),$AL$9<>"",B13=$AL$9)',
        fill=PERIOD_COLORS["ready"],
        font_color="FFFFFF",
        bold=True,
        stop=True,
    )
    # Holiday/hiatus override: selected-location holidays affecting Production are pink.
    _add_cf(
        ws,
        [combined],
        '=AND(ISNUMBER(B13),WEEKDAY(B13,2)<=5,B13>=$AL$5,B13<=$AM$5,COUNTIF($AQ$3:$AQ$80,B13)>0)',
        fill=PERIOD_COLORS["hiatus"],
        stop=True,
    )
    # Period colors. Visual bands are weekday-only, matching the reference workbook.
    helper_rows = {"rd": 2, "pre": 3, "travel": 4, "production": 5, "hiatus": 6, "post": 7, "print_ship": 8}
    for key in ["rd", "pre", "travel", "production", "hiatus", "post", "print_ship"]:
        r = helper_rows[key]
        _add_cf(
            ws,
            [combined],
            f'=AND(ISNUMBER(B13),WEEKDAY(B13,2)<=5,$AP${r}=TRUE,$AL${r}<>"",$AM${r}<>"",B13>=$AL${r},B13<=$AM${r})',
            fill=PERIOD_COLORS[key],
        )
    # Light gray font on weekends, similar to the original template behavior.
    _add_cf(ws, [combined], '=OR(WEEKDAY(B13,1)=1,WEEKDAY(B13,1)=7)', font_color="808080")


def _week_label_formula(row_range: str) -> str:
    start_expr = f"AGGREGATE(15,6,{row_range}/({row_range}<>\"\"),1)"
    end_expr = f"AGGREGATE(14,6,{row_range}/({row_range}<>\"\"),1)"
    def back_week(helper_row: int) -> str:
        return f'MAX(1,$AN${helper_row}-INT((NETWORKDAYS($AL${helper_row},MAX({start_expr},$AL${helper_row}))-1)/5))'
    def fwd_week(helper_row: int, holidays: bool = False) -> str:
        harg = ',Holidays!$A$4:$A$200' if holidays else ''
        return f'INT((NETWORKDAYS($AL${helper_row},MAX({start_expr},$AL${helper_row}){harg})-1)/5)+1'
    # Week numbers are based on Monday-Friday workweeks. Production also excludes selected-location holidays.
    return (
        f'=IFERROR(IF(AND({end_expr}>=$AL$2,{start_expr}<=$AM$2),{back_week(2)},'
        f'IF(AND({end_expr}>=$AL$3,{start_expr}<=$AM$3),{back_week(3)},'
        f'IF(AND({end_expr}>=$AL$4,{start_expr}<=$AM$4),{back_week(4)},'
        f'IF(AND({end_expr}>=$AL$5,{start_expr}<=$AM$5),{fwd_week(5, True)},'
        f'IF(AND({end_expr}>=$AL$6,{start_expr}<=$AM$6),"H",'
        f'IF(AND({end_expr}>=$AL$7,{start_expr}<=$AM$7),{fwd_week(7)},'
        f'IF(AND({end_expr}>=$AL$8,{start_expr}<=$AM$8),{fwd_week(8)},""))))))),"")'
    )

def _populate_week_labels(ws: Worksheet) -> None:
    for row_start in DATE_ROW_STARTS:
        for gutter_col, start_col, end_col in DATE_BLOCKS:
            for r in range(row_start, row_start + 6):
                row_range = f"{start_col}{r}:{end_col}{r}"
                c = ws[f"{gutter_col}{r}"]
                c.value = _week_label_formula(row_range)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = _font("000000", True, 9)


def _date_box_periods(schedule: Dict[str, Any]) -> List[str]:
    present = {p.get("key") for p in schedule.get("periods", [])}
    ordered: List[str] = []
    for key in ["rd", "pre", "travel", "production", "hiatus", "post", "print_ship"]:
        if key in present:
            ordered.append(key)
    ordered.append("ready")
    return ordered[:8]


def _schedule_data_row_for_key(key: str) -> int:
    return {"rd": 3, "pre": 4, "travel": 5, "production": 6, "hiatus": 7, "post": 8, "print_ship": 9, "ready": 10}[key]


def _unmerge_intersecting(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for rng in list(ws.merged_cells.ranges):
        if rng.max_row >= min_row and rng.min_row <= max_row and rng.max_col >= min_col and rng.min_col <= max_col:
            ws.unmerge_cells(str(rng))


def _populate_lower_box(ws: Worksheet, schedule: Dict[str, Any]) -> None:
    """Build the bottom Dates box and the adjacent color palette legend.

    User-specified target:
      - Dates label above the box at B37.
      - Dates box outer border spans B38:P44 exactly.
      - Left/right borders are B38:B44 and P38:P44.
      - Legend swatches are two cells wide, R:S, starting at R38:S38.
      - Legend labels sit to the right as plain black text.
    """
    _unmerge_intersecting(ws, 36, 50, 2, 33)  # B:AG

    # Clear the lower area completely so no old merged cells, fills, or borders
    # survive into either the XLSX or the PDF conversion.
    for r in range(36, 51):
        ws.row_dimensions[r].height = 18
        for c in range(2, 34):  # B:AG
            cell = ws.cell(r, c)
            if not isinstance(cell, MergedCell):
                cell.value = None
                cell.fill = PatternFill(fill_type=None)
                cell.font = _font("000000", False, 10)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = Border()

    # Do not change column widths here. The same global column widths serve
    # both the calendar grid and the lower Dates/Legend area so the right-side
    # months never become wider or narrower after export.
    _set_reference_column_widths(ws)

    # Dates label above the fixed box.
    ws["B37"] = "Dates:"
    ws["B37"].font = _font("000000", True, 10)
    ws["B37"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[37].height = 15

    first_row = 38
    last_row = 44
    text_start, text_end = 2, 11   # B:K
    metric_start, metric_end = 12, 16  # L:P
    thick = Side(style="medium", color="000000")

    present = {p.get("key") for p in schedule.get("periods", [])}

    # Build the narrative rows. R&D and Travel/Prep are included only when used.
    # Keep the common no-hiatus/pre-prod layout aligned with the reference crop:
    # Pre row, blank spacer, Production row, Post row, Print & Ship, Ready.
    early_keys = [key for key in ["rd", "pre", "travel"] if key in present]
    has_hiatus = "hiatus" in present
    has_post = "post" in present
    has_print = "print_ship" in present
    has_production = "production" in present

    row_for_key: Dict[str, int] = {}
    spacer_rows: set[int] = set()
    production_metric_continuation_row: Optional[int] = None

    if not has_hiatus and early_keys == ["pre"] and has_production and has_post and has_print:
        row_for_key = {"pre": 38, "production": 40, "post": 42, "print_ship": 43, "ready": 44}
        spacer_rows.add(39)
        production_metric_continuation_row = 41
    else:
        current = first_row
        keys = early_keys + (["production"] if has_production else []) + (["hiatus"] if has_hiatus else []) + (["post"] if has_post else []) + (["print_ship"] if has_print else []) + ["ready"]
        for key in keys:
            if current > last_row:
                # Always preserve Ready for Release in the final visible row.
                if key == "ready":
                    row_for_key[key] = last_row
                break
            row_for_key[key] = current
            # Production's metric is two lines, so give it more vertical space,
            # but keep it within a single row so the box remains B38:P44.
            current += 1

    def merge_line(row: int) -> Tuple[Any, Any]:
        text_cell = ws.cell(row, text_start)
        metric_cell = ws.cell(row, metric_start)
        text_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        metric_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        return text_cell, metric_cell

    # Apply row heights first so they stay stable after merging.
    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = 20
    for r in spacer_rows:
        ws.row_dimensions[r].height = 18
    if "production" in row_for_key and production_metric_continuation_row is None:
        ws.row_dimensions[row_for_key["production"]].height = 34

    # Merge every row of the fixed Dates box into a left narrative zone and a
    # right metric zone, including blank spacer rows, so the box behaves like
    # the reference rather than a grid.
    for r in range(first_row, last_row + 1):
        ws.merge_cells(start_row=r, start_column=text_start, end_row=r, end_column=text_end)
        ws.merge_cells(start_row=r, start_column=metric_start, end_row=r, end_column=metric_end)
        merge_line(r)

    for key, row in row_for_key.items():
        if row < first_row or row > last_row:
            continue
        text_cell, metric_cell = merge_line(row)
        sd_row = _schedule_data_row_for_key(key)
        text_cell.value = f"='Schedule Data'!$J${sd_row}"
        text_cell.font = _font("000000", False, 10)
        metric_cell.value = f"='Schedule Data'!$K${sd_row}" if key not in ("ready", "hiatus", "print_ship") else ""
        metric_cell.font = _font("000000", True, 10)
        metric_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if key == "production":
            if production_metric_continuation_row:
                metric_cell.value = f"=IFERROR(LEFT('Schedule Data'!$K${sd_row},FIND(CHAR(10),'Schedule Data'!$K${sd_row})-1),'Schedule Data'!$K${sd_row})"
                continuation_cell = ws.cell(production_metric_continuation_row, metric_start)
                continuation_cell.value = f"=IFERROR(MID('Schedule Data'!$K${sd_row},FIND(CHAR(10),'Schedule Data'!$K${sd_row})+1,99),\"\")"
                continuation_cell.font = _font("000000", True, 10)
                continuation_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                metric_cell.value = f"='Schedule Data'!$K${sd_row}"
        elif key == "print_ship":
            text_cell.font = Font(name="Arial", size=10, bold=False, italic=True, color="000000")
        elif key == "ready":
            text_cell.font = _font("000000", True, 10)
        elif key == "hiatus":
            # Hiatus has no duration metric in the reference-style narrative.
            metric_cell.value = ""

    # Thick outer border around B38:P44 exactly. No internal borders.
    for r in range(first_row, last_row + 1):
        for c in range(2, 17):  # B:P
            cell = ws.cell(r, c)
            cell.border = Border(
                left=thick if c == 2 else Side(style=None),
                right=thick if c == 16 else Side(style=None),
                top=thick if r == first_row else Side(style=None),
                bottom=thick if r == last_row else Side(style=None),
            )

    # Notes label directly below the fixed box.
    notes_row = last_row + 1
    ws[f"B{notes_row}"] = "Notes:"
    ws[f"B{notes_row}"].font = _font("000000", True, 10)
    ws[f"B{notes_row}"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[notes_row].height = 16

    # Adjacent color palette legend: swatches are exactly R:S, one row per item.
    prod = next((p for p in schedule.get("periods", []) if p.get("key") == "production"), None)
    skipped = ((prod or {}).get("metric") or {}).get("skippedHolidays", 0)
    legend_items: List[Tuple[str, str]] = []
    for key in ["rd", "pre", "travel"]:
        if key in present:
            legend_items.append((key, PERIOD_LABELS[key]))
    if "production" in present:
        legend_items.append(("production", PERIOD_LABELS["production"]))
        legend_items.append(("hiatus", "Holidays / Hiatus"))
    elif skipped or "hiatus" in present:
        legend_items.append(("hiatus", "Holidays / Hiatus"))
    for key in ["post", "print_ship"]:
        if key in present:
            legend_items.append((key, PERIOD_LABELS[key]))

    legend_row = first_row
    for key, label in legend_items[:7]:
        if legend_row > last_row:
            break
        # Swatch is exactly two cells wide: R:S. Do not merge the swatch cells;
        # applying the outer border to both cells makes the thick box display
        # consistently in Excel, Numbers, and PDF conversion.
        ws.merge_cells(start_row=legend_row, start_column=21, end_row=legend_row, end_column=33)  # U:AG label
        left_swatch = ws.cell(legend_row, 18)
        right_swatch = ws.cell(legend_row, 19)
        label_cell = ws.cell(legend_row, 21)
        for swatch_cell, side in ((left_swatch, "left"), (right_swatch, "right")):
            swatch_cell.fill = _fill(PERIOD_COLORS[key])
            swatch_cell.alignment = Alignment(horizontal="center", vertical="center")
            swatch_cell.border = Border(
                left=thick if side == "left" else Side(style=None),
                right=thick if side == "right" else Side(style=None),
                top=thick,
                bottom=thick,
            )
        label_cell.value = label
        label_cell.font = _font("000000", False, 10)
        label_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        label_cell.fill = PatternFill(fill_type=None)
        label_cell.border = Border()
        legend_row += 1


def _finalize_year_sheet(ws: Worksheet, year: int, schedule: Dict[str, Any]) -> None:
    _prepare_year_sheet(ws, year)
    _populate_year_headers(ws, year)
    _populate_local_helpers(ws, year, schedule)
    _populate_week_labels(ws)
    _populate_lower_box(ws, schedule)
    _apply_calendar_conditional_formatting(ws)


def _ensure_year_sheets(wb, years: List[int]) -> List[Worksheet]:
    # Use the 2026 tab as the visual master because it contains the reference layout.
    if "2026" not in wb.sheetnames:
        raise ValueError("Reference template must include a 2026 sheet.")
    master = wb["2026"]
    sheets: List[Worksheet] = []
    existing_by_year = {int(ws.title): ws for ws in wb.worksheets if ws.title.isdigit()}
    for year in years:
        if year in existing_by_year:
            ws = existing_by_year[year]
        else:
            ws = wb.copy_worksheet(master)
            ws.title = str(year)
        sheets.append(ws)
    # Remove year sheets not touched by the schedule.
    for ws in list(wb.worksheets):
        if ws.title.isdigit() and int(ws.title) not in years:
            wb.remove(ws)
    # Order year sheets after Inputs.
    for year in years:
        ws = wb[str(year)]
        wb._sheets.remove(ws)
        wb._sheets.append(ws)
    return [wb[str(y)] for y in years]


def create_workbook(payload: Dict[str, Any], out_path: str | Path) -> Path:
    schedule = calculate_schedule(payload)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Reference Excel template missing: {TEMPLATE_PATH}")

    years = schedule.get("years") or [date.today().year]
    years = sorted(set(int(y) for y in years))

    wb = load_workbook(TEMPLATE_PATH)
    year_sheets = _ensure_year_sheets(wb, years)
    _clean_branding(wb)
    _make_inputs_sheet(wb, schedule, payload)
    _make_holidays_sheet(wb, schedule, years)
    _make_schedule_data_sheet(wb, schedule, payload)

    # Hide helper sheets by default; Inputs stays visible so exported files remain editable.
    wb["Schedule Data"].sheet_state = "hidden"
    wb["Holidays"].sheet_state = "hidden"

    for y in years:
        _finalize_year_sheet(wb[str(y)], y, schedule)

    # Move Inputs to the front, year sheets next, hidden helpers after.
    ordered = [wb["Inputs"]] + [wb[str(y)] for y in years] + [wb["Schedule Data"], wb["Holidays"]]
    wb._sheets = ordered
    wb.active = 1 if years else 0

    # Force full recalculation when opened in Excel.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    # Final scan for banned phrases.
    _clean_branding(wb)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    sample = {
        "projectTitle": "THE COMEBACKER",
        "productionLocation": "US",
        "anchorMode": "production",
        "periods": {
            "rd": {"start": "", "weeks": 0},
            "pre": {"start": "", "weeks": 12},
            "travel": {"start": "", "weeks": 0},
            "production": {"start": "10/05/26", "days": 40},
            "hiatus": {"start": "", "end": ""},
            "post": {"start": "", "weeks": 24},
            "print_ship": {"start": "", "weeks": 4},
            "ready": {"date": "06/18/27"},
        },
    }
    create_workbook(sample, APP_DIR / "exports" / "sample_spe_block_calendar.xlsx")
